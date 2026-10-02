import sqlite3
import openpyxl

print("=== RECONCILIATION ANALYSIS ===")

# 1. Inspect telecom_portal.db
conn = sqlite3.connect('telecom_portal.db')
cursor = conn.cursor()

cursor.execute("SELECT COUNT(*) FROM collo_records")
db_collo_count = cursor.fetchone()[0]

cursor.execute("SELECT status, COUNT(*) FROM collo_records GROUP BY status")
db_collo_status = dict(cursor.fetchall())

cursor.execute("SELECT SUM(rev_sewa_tahun), SUM(biaya_otc), SUM(biaya_sewa_tahun), SUM(margin_rupiah) FROM collo_records WHERE status = 'ACTIVE'")
db_collo_fin = cursor.fetchone()

cursor.execute("SELECT COUNT(*) FROM sitac_records")
db_sitac_count = cursor.fetchone()[0]

cursor.execute("SELECT progress, COUNT(*) FROM sitac_records GROUP BY progress")
db_sitac_progress = dict(cursor.fetchall())

cursor.execute("SELECT SUM(biaya_permintaan_awal), SUM(biaya_final), SUM(efisiensi_rupiah) FROM sitac_records")
db_sitac_fin = cursor.fetchone()

cursor.execute("SELECT COUNT(*) FROM gangguan_records")
db_gangguan_count = cursor.fetchone()[0]

cursor.execute("SELECT is_selesai, COUNT(*) FROM gangguan_records GROUP BY is_selesai")
db_gangguan_status = dict(cursor.fetchall())

cursor.execute("SELECT SUM(biaya_gangguan) FROM gangguan_records")
db_gangguan_cost = cursor.fetchone()[0]

print("\n--- DATABASE CURRENT STATE (telecom_portal.db) ---")
print(f"Collocation Records: {db_collo_count}")
print(f"  Status breakdown: {db_collo_status}")
print(f"  ACTIVE Financials: Rev={db_collo_fin[0]:,.0f}, OTC={db_collo_fin[1]:,.0f}, Biaya={db_collo_fin[2]:,.0f}, Margin={db_collo_fin[3]:,.0f}")

print(f"\nSITAC Records: {db_sitac_count}")
print(f"  Progress breakdown: {db_sitac_progress}")
print(f"  Financials: Awal={db_sitac_fin[0]:,.0f}, Final={db_sitac_fin[1]:,.0f}, Hemat={db_sitac_fin[2]:,.0f}")

print(f"\nGangguan Records: {db_gangguan_count}")
print(f"  Status breakdown (0=Open, 1=Selesai): {db_gangguan_status}")
print(f"  Total Cost: Rp {db_gangguan_cost:,.0f}")

# 2. Inspect 'Report Sitac  2026 FIX.xlsx'
print("\n--- INSPECTING Report Sitac  2026 FIX.xlsx ---")
wb_sitac = openpyxl.load_workbook('Report Sitac  2026 FIX.xlsx', data_only=True, read_only=True)
print("Sheets:", wb_sitac.sheetnames)

# Master RAW data (SITAC)
sheet_raw = wb_sitac['Master RAW data']
raw_rows = 0
raw_progress = {}
raw_awal = 0
raw_final = 0
for r_idx, row in enumerate(sheet_raw.iter_rows(values_only=True)):
    if r_idx == 0: continue
    no_pa = row[1]
    pelanggan = row[2]
    # Check if row is not completely empty
    if not no_pa and not pelanggan and not row[4]:
        continue
    raw_rows += 1
    prog = str(row[13] or 'Ongoing').strip().title()
    raw_progress[prog] = raw_progress.get(prog, 0) + 1
    if isinstance(row[9], (int, float)): raw_awal += row[9]
    if isinstance(row[11], (int, float)): raw_final += row[11]

print(f"Excel SITAC (Master RAW data) total rows: {raw_rows}")
print(f"  Progress breakdown: {raw_progress}")
print(f"  Sum Biaya Awal: {raw_awal:,.0f}")
print(f"  Sum Biaya Final: {raw_final:,.0f}")

# Gangguan sheet
if 'GANGGUAN' in wb_sitac.sheetnames:
    sheet_g = wb_sitac['GANGGUAN']
    g_rows = 0
    g_status = {}
    g_cost = 0
    headers_g = None
    for r_idx, row in enumerate(sheet_g.iter_rows(values_only=True)):
        if r_idx == 0:
            headers_g = row
            continue
        no_t = row[1] if len(row) > 1 else None
        if not no_t and not (len(row) > 2 and row[2]):
            continue
        g_rows += 1
        # Try to find cost and status
        # Let's inspect headers on first print
    print(f"Excel GANGGUAN sheet total rows: {g_rows}")
    if headers_g:
        print("  Headers GANGGUAN:", [f"{idx}:{h}" for idx, h in enumerate(headers_g) if h is not None][:15])

# 3. Inspect 'Rekap Collo x Interkoneksi 2024 V2 ROM 22 Nov(1).xlsx'
print("\n--- INSPECTING Rekap Collo x Interkoneksi 2024 V2 ROM 22 Nov(1).xlsx ---")
wb_collo = openpyxl.load_workbook('Rekap Collo x Interkoneksi 2024 V2 ROM 22 Nov(1).xlsx', data_only=True, read_only=True)
print("Sheets:", wb_collo.sheetnames)

for sname in wb_collo.sheetnames:
    s = wb_collo[sname]
    r_count = 0
    for r_idx, r in enumerate(s.iter_rows(values_only=True)):
        if any(c is not None for c in r):
            r_count += 1
    print(f"  Sheet '{sname}': {r_count} rows with content")
