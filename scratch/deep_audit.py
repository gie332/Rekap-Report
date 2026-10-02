import sqlite3
import openpyxl

conn = sqlite3.connect('telecom_portal.db')
cursor = conn.cursor()

# ==========================================
# 1. SITAC DETAILED ROW-BY-ROW COMPARISON
# ==========================================
wb_s = openpyxl.load_workbook('Report Sitac  2026 FIX.xlsx', data_only=True, read_only=True)
sheet_s = wb_s['Master RAW data']

excel_sitac = []
for idx, r in enumerate(sheet_s.iter_rows(values_only=True)):
    if idx == 0: continue
    no_pa = str(r[1] or '').strip()
    pelanggan = str(r[2] or '').strip()
    if not no_pa and not pelanggan: continue
    pic = str(r[4] or '').strip()
    b_awal = float(r[9]) if isinstance(r[9], (int, float)) else 0.0
    b_final = float(r[11]) if isinstance(r[11], (int, float)) else 0.0
    prog = str(r[13] or 'Ongoing').strip().title()
    excel_sitac.append({
        'row': idx + 1,
        'no_pa': no_pa,
        'pelanggan': pelanggan,
        'pic': pic,
        'b_awal': b_awal,
        'b_final': b_final,
        'prog': prog
    })

cursor.execute("SELECT id, no_pa, pelanggan, pic_perijinan, biaya_permintaan_awal, biaya_final, efisiensi_rupiah, progress FROM sitac_records ORDER BY id ASC")
db_sitac = [dict(zip(['id', 'no_pa', 'pelanggan', 'pic', 'b_awal', 'b_final', 'efisiensi', 'prog'], row)) for row in cursor.fetchall()]

print(f"Total SITAC in Excel: {len(excel_sitac)}, in DB: {len(db_sitac)}")

# Compare financial differences
diff_count = 0
for idx in range(min(len(excel_sitac), len(db_sitac))):
    ex = excel_sitac[idx]
    db = db_sitac[idx]
    if abs(ex['b_awal'] - db['b_awal']) > 1 or abs(ex['b_final'] - db['b_final']) > 1 or ex['prog'].lower() != db['prog'].lower():
        diff_count += 1
        if diff_count <= 10:
            print(f"Diff at #{idx+1} (ID {db['id']}): Excel(Awal={ex['b_awal']:,.0f}, Fin={ex['b_final']:,.0f}, Prog={ex['prog']}) vs DB(Awal={db['b_awal']:,.0f}, Fin={db['b_final']:,.0f}, Prog={db['prog']}) [PA: {ex['no_pa']}]")

print(f"Total SITAC rows with discrepancies: {diff_count} / {len(excel_sitac)}")

# ==========================================
# 2. GANGGUAN COMPARISON
# ==========================================
sheet_g = wb_s['GANGGUAN']
excel_gangguan = []
for idx, r in enumerate(sheet_g.iter_rows(values_only=True)):
    if idx == 0: continue
    no_t = str(r[3] or r[1] or '').strip()
    term = str(r[4] or '').strip()
    if not no_t and not term: continue
    # Let's inspect columns for biaya and status
    b_awal = float(r[9]) if len(r) > 9 and isinstance(r[9], (int, float)) else 0.0
    b_final = float(r[10]) if len(r) > 10 and isinstance(r[10], (int, float)) else 0.0
    excel_gangguan.append({
        'row': idx + 1,
        'tiket': no_t,
        'terminating': term,
        'b_awal': b_awal,
        'b_final': b_final
    })

cursor.execute("SELECT id, no_tiket, terminating, biaya_gangguan, is_selesai FROM gangguan_records ORDER BY id ASC")
db_gangguan = [dict(zip(['id', 'tiket', 'terminating', 'biaya', 'selesai'], row)) for row in cursor.fetchall()]

print(f"\nTotal GANGGUAN in Excel: {len(excel_gangguan)}, in DB: {len(db_gangguan)}")

# ==========================================
# 3. COLLO STATUS & TOTALS IN EXCEL
# ==========================================
wb_c = openpyxl.load_workbook('Rekap Collo x Interkoneksi 2024 V2 ROM 22 Nov(1).xlsx', data_only=True, read_only=True)
sheet_c = wb_c['Raw Data']
collo_headers = None
c_status = {}
c_rev_act = 0
c_biaya_act = 0
c_otc_act = 0
total_raw_rows = 0

for idx, r in enumerate(sheet_c.iter_rows(values_only=True)):
    if idx == 0:
        collo_headers = r
        continue
    if not any(x is not None for x in r): continue
    total_raw_rows += 1
    # Find status column
    # Typically status is around col 25
    st = str(r[24] if len(r) > 24 else '').strip().upper()
    c_status[st] = c_status.get(st, 0) + 1
    rev = float(r[13]) if len(r) > 13 and isinstance(r[13], (int, float)) else 0.0
    otc = float(r[12]) if len(r) > 12 and isinstance(r[12], (int, float)) else 0.0
    biaya = float(r[16]) if len(r) > 16 and isinstance(r[16], (int, float)) else 0.0
    if st == 'ACTIVE':
        c_rev_act += rev
        c_otc_act += otc
        c_biaya_act += biaya

print(f"\nTotal Collocation in Excel 'Raw Data': {total_raw_rows}")
print(f"  Status breakdown in Excel: {c_status}")
print(f"  ACTIVE Financials in Excel: Rev={c_rev_act:,.0f}, OTC={c_otc_act:,.0f}, Biaya={c_biaya_act:,.0f}")
