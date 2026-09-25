#!/usr/bin/env python3
"""
extract_sitac_data.py
Pipeline pembersihan & ekstraksi komprehensif dari file Excel:
'Report Sitac  2026 FIX.xlsx'

Sheets yang diproses:
1. 'Master RAW data': Data utama Project Authorization (PA) SITAC & Perizinan
2. 'Rekap Efesiensi Biaya Perijinan': Ringkasan efisiensi biaya tahun 2024-2026
3. 'GANGGUAN': Data tiket dan insiden gangguan jaringan/perizinan di lapangan

Output:
- 'sitac_data.json'      : Format JSON terstruktur lengkap & teroptimasi
- 'sitac_data.js'        : Format JS Standalone (window.SITAC_RAW_DATA) untuk offline/file:// protocol
- 'sitac_database.db'    : Database lokal SQLite dengan tabel, indeks, dan views
"""

import os
import re
import json
import sqlite3
import datetime
import openpyxl

EXCEL_FILE = "Report Sitac  2026 FIX.xlsx"
OUTPUT_JSON = "sitac_data.json"
OUTPUT_JS = "sitac_data.js"
OUTPUT_DB = "sitac_database.db"

# Anchor geografis Jabodetabek & Jawa Barat untuk geocoding alamat terminating
GEO_ANCHORS = {
    'tanah abang': (-6.1855, 106.8155),
    'kebon sirih': (-6.1837, 106.8290),
    'menteng': (-6.1950, 106.8320),
    'kemayoran': (-6.1550, 106.8520),
    'cideng': (-6.1730, 106.8110),
    'merdeka barat': (-6.1760, 106.8230),
    'sudirman': (-6.2150, 106.8220),
    'kuningan': (-6.2280, 106.8310),
    'setiabudi': (-6.2110, 106.8280),
    'duren tiga': (-6.2520, 106.8390),
    'tebet': (-6.2300, 106.8520),
    'ancol': (-6.1260, 106.8320),
    'pluit': (-6.1180, 106.7910),
    'sunter': (-6.1420, 106.8710),
    'marunda': (-6.1050, 106.9650),
    'kelapa gading': (-6.1580, 106.9080),
    'daan mogot': (-6.1580, 106.7350),
    'cengkareng': (-6.1450, 106.7280),
    'grogol': (-6.1660, 106.7890),
    'halim': (-6.2620, 106.8910),
    'cawang': (-6.2480, 106.8680),
    'cakung': (-6.1880, 106.9510),
    'buaran': (-6.2200, 106.9180),
    'bekasi': (-6.2380, 106.9750),
    'harapan indah': (-6.1830, 106.9820),
    'cikarang': (-6.2610, 107.1520),
    'depok': (-6.3950, 106.8230),
    'limo': (-6.3680, 106.7720),
    'tangerang': (-6.1780, 106.6310),
    'bsd': (-6.3010, 106.6520),
    'serpong': (-6.3210, 106.6690),
    'greenlake': (-6.1850, 106.7020),
    'bogor': (-6.5950, 106.7900),
    'bojong': (-6.4950, 106.7950),
    'serang': (-6.1100, 106.1500),
    'karawang': (-6.3150, 107.2950),
    'jakarta': (-6.2000, 106.8200)
}

def parse_currency(val):
    """
    Mengonversi berbagai representasi biaya (angka, string dengan titik,
    simbol 'Rp', spasi, atau teks non-moneter) menjadi nilai integer rupiah dan catatan.
    """
    if val is None:
        return 0, ""
    
    if isinstance(val, (int, float)):
        return int(round(val)), ""
    
    val_str = str(val).strip()
    if not val_str or val_str in ("-", "None", "0"):
        return 0, ""
    
    # Deteksi teks non-moneter / in-kind seperti 'Bak sampah fiber 660 liter'
    in_kind_keywords = ['sampah', 'liter', 'meter', 'kabel', 'fiber', 'unit', 'buah', 'semen', 'cat', 'wifi']
    if any(k in val_str.lower() for k in in_kind_keywords):
        return 0, val_str

    if not re.search(r'\d', val_str):
        return 0, val_str

    # Bersihkan simbol Rp, spasi, dan karakter non-angka kecuali titik/koma
    cleaned = re.sub(r'[^\d.,]', '', val_str)
    
    if '.' in cleaned and ',' not in cleaned:
        parts = cleaned.split('.')
        if len(parts) > 1 and all(len(p) == 3 for p in parts[1:]):
            cleaned = cleaned.replace('.', '')
        else:
            cleaned = cleaned.replace('.', '')
    elif ',' in cleaned and '.' in cleaned:
        cleaned = cleaned.replace('.', '').replace(',', '.')
    elif ',' in cleaned:
        cleaned = cleaned.replace(',', '.')

    try:
        num = float(cleaned)
        return int(round(num)), (val_str if " " in val_str and not val_str.startswith("Rp") else "")
    except ValueError:
        return 0, val_str

def parse_date(val):
    """Mengonversi nilai cell tanggal menjadi string format YYYY-MM-DD."""
    if val is None:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("none", "null", "-", "0000-00-00"):
        return None
    
    # Coba format ISO YYYY-MM-DD
    match = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', val_str)
    if match:
        y, m, d = match.groups()
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    
    # Coba format DD/MM/YYYY atau DD-MM-YYYY
    match = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})', val_str)
    if match:
        d, m, y = match.groups()
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
        
    return None

def normalize_status(val):
    """Menormalisasi nilai progress status."""
    if not val:
        return "Finish"
    s = str(val).strip().lower()
    if "finish" in s or "selesai" in s:
        return "Finish"
    elif "cancel" in s or "batal" in s:
        return "Cancel"
    elif "ongoing" in s or "proses" in s or "jalan" in s:
        return "Ongoing"
    elif "hold" in s or "tunda" in s:
        return "Hold"
    return s.capitalize()

def parse_int_safe(val, default=None):
    """Konversi aman ke integer."""
    if val is None:
        return default
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return default

def clean_text(val):
    """Pembersihan teks standar."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ("none", "null", "-"):
        return ""
    return s

def resolve_coordinates(coord_text, address_text, record_id):
    """
    Mengekstrak atau mengestimasi koordinat geografis (lat, lng)
    berdasarkan teks koordinat atau nama wilayah/jalan di Jabodetabek.
    """
    coord_regex = re.compile(r'(-?\d{1,2}\.\d+)\s*,\s*(\d{2,3}\.\d+)')
    
    # 1. Cek dari teks koordinat langsung
    if coord_text:
        match = coord_regex.search(coord_text)
        if match:
            try:
                lat = float(match.group(1))
                lng = float(match.group(2))
                return round(lat, 6), round(lng, 6), "EXACT"
            except ValueError:
                pass

    # 2. Cek dari teks alamat terminating
    if address_text:
        match = coord_regex.search(address_text)
        if match:
            try:
                lat = float(match.group(1))
                lng = float(match.group(2))
                return round(lat, 6), round(lng, 6), "EXACT"
            except ValueError:
                pass

        # 3. Estimasi berdasarkan kata kunci wilayah Jabodetabek
        lower_addr = address_text.lower()
        for key, (base_lat, base_lng) in GEO_ANCHORS.items():
            if key in lower_addr:
                # Berikan sedikit variasi realistis berdasarkan record_id agar marker tidak menumpuk tepat di titik yang sama
                offset_lat = ((record_id * 17) % 31 - 15) * 0.0018
                offset_lng = ((record_id * 23) % 31 - 15) * 0.0018
                return round(base_lat + offset_lat, 6), round(base_lng + offset_lng, 6), "GEOCODED"

    # Default fallback di area Jakarta Pusat dengan sebaran acak tertata
    offset_lat = ((record_id * 13) % 41 - 20) * 0.003
    offset_lng = ((record_id * 19) % 41 - 20) * 0.003
    return round(-6.2088 + offset_lat, 6), round(106.8456 + offset_lng, 6), "AREA_DEFAULT"

# ==============================================================================
# 1. EXTRACTION: MASTER RAW DATA
# ==============================================================================
def extract_master_raw_data(wb):
    print("[INFO] Mengekstrak sheet: Master RAW data...")
    ws = wb["Master RAW data"]
    rows = list(ws.iter_rows(values_only=True))
    data_rows = rows[1:]

    records = []
    for row in data_rows:
        if not any(row):
            continue
        
        no_val = row[0] if len(row) > 0 else None
        no_pa = clean_text(row[1] if len(row) > 1 else None)
        pelanggan = clean_text(row[2] if len(row) > 2 else None)
        
        # Abaikan baris kosong
        if no_val is None and not no_pa and not pelanggan:
            continue
            
        seq_id = len(records) + 1
        excel_no = parse_int_safe(no_val, seq_id)
        ptl = clean_text(row[3] if len(row) > 3 else "")
        perijinan = clean_text(row[4] if len(row) > 4 else "")
        tiering = clean_text(row[5] if len(row) > 5 else "")
        poin = parse_int_safe(row[6] if len(row) > 6 else None)
        terminating = clean_text(row[7] if len(row) > 7 else "")
        koordinat_raw = clean_text(row[8] if len(row) > 8 else "")
        
        # Resolusi koordinat untuk peta WebGIS
        lat, lng, coord_type = resolve_coordinates(koordinat_raw, terminating, seq_id)
        
        # Biaya Permintaan Awal & Final
        raw_biaya_awal = row[9] if len(row) > 9 else None
        biaya_awal, catatan_biaya_awal = parse_currency(raw_biaya_awal)
        
        sewa_otc = clean_text(row[10] if len(row) > 10 else "")
        if not sewa_otc:
            sewa_otc = "OTC" if (biaya_awal > 0 or (len(row) > 11 and parse_currency(row[11])[0] > 0)) else "-"
            
        raw_biaya_final = row[11] if len(row) > 11 else None
        biaya_final, _ = parse_currency(raw_biaya_final)
        
        raw_biaya_sewa = row[12] if len(row) > 12 else None
        biaya_sewa_bulan, _ = parse_currency(raw_biaya_sewa)
        
        progress = normalize_status(row[13] if len(row) > 13 else "Finish")
        
        date_pa = parse_date(row[14] if len(row) > 14 else None)
        date_dispos = parse_date(row[15] if len(row) > 15 else None)
        date_close = parse_date(row[18] if len(row) > 18 else None)
        
        aging_dispos = parse_int_safe(row[16] if len(row) > 16 else None)
        stopclock = parse_int_safe(row[17] if len(row) > 17 else None)
        aging_close = parse_int_safe(row[20] if len(row) > 20 else None)
        aging_berjalan = parse_int_safe(row[21] if len(row) > 21 else None)
        
        if aging_close is None and date_dispos and date_close:
            try:
                d1 = datetime.datetime.strptime(date_dispos, "%Y-%m-%d")
                d2 = datetime.datetime.strptime(date_close, "%Y-%m-%d")
                aging_close = max(0, (d2 - d1).days)
            except Exception:
                pass
                
        durasi_hari = aging_close if progress == "Finish" else (aging_berjalan if aging_berjalan is not None else 0)
        
        # Klasifikasi Aging Indicator:
        # Hijau: <= 7 hari (Aman)
        # Kuning: 8 - 14 hari (> 7 hari)
        # Merah: > 14 hari (> 14 hari)
        if durasi_hari is None or durasi_hari < 0:
            aging_category = "green"
            aging_label = "Aman"
        elif durasi_hari <= 7:
            aging_category = "green"
            aging_label = "Aman (<=7 Hari)"
        elif durasi_hari <= 14:
            aging_category = "yellow"
            aging_label = "Perhatian (8-14 Hari)"
        else:
            aging_category = "red"
            aging_label = "Kritis (>14 Hari)"

        update_pekerjaan = clean_text(row[22] if len(row) > 22 else "")
        bulan = parse_int_safe(row[23] if len(row) > 23 else None)
        tahun = parse_int_safe(row[24] if len(row) > 24 else None)
        bulan_finish = parse_int_safe(row[25] if len(row) > 25 else None)
        tahun_finish = parse_int_safe(row[26] if len(row) > 26 else None)
        
        if not tahun and date_dispos:
            tahun = int(date_dispos.split("-")[0])
            bulan = int(date_dispos.split("-")[1])
        elif not tahun and date_close:
            tahun = int(date_close.split("-")[0])
            bulan = int(date_close.split("-")[1])
        if not tahun:
            tahun = 2025
            
        # Efisiensi negosiasi biaya
        if biaya_awal > 0:
            efisiensi_rupiah = max(0, biaya_awal - biaya_final)
            efisiensi_persen = round((efisiensi_rupiah / biaya_awal) * 100, 2)
        else:
            efisiensi_rupiah = 0
            efisiensi_persen = 0.0

        is_berbiaya = 1 if (biaya_final > 0 or biaya_awal > 0) else 0

        records.append({
            "id": seq_id,
            "excel_no": excel_no,
            "no_pa": no_pa,
            "pelanggan": pelanggan,
            "ptl": ptl,
            "pic_perijinan": perijinan,
            "tiering": tiering,
            "poin": poin,
            "terminating": terminating,
            "koordinat_raw": koordinat_raw,
            "latitude": lat,
            "longitude": lng,
            "coord_type": coord_type,
            "biaya_permintaan_awal": biaya_awal,
            "catatan_biaya_awal": catatan_biaya_awal,
            "sewa_otc": sewa_otc,
            "biaya_final": biaya_final,
            "biaya_sewa_bulan": biaya_sewa_bulan,
            "efisiensi_rupiah": efisiensi_rupiah,
            "efisiensi_persen": efisiensi_persen,
            "is_berbiaya": is_berbiaya,
            "progress": progress,
            "date_pa": date_pa,
            "date_dispos": date_dispos,
            "date_close": date_close,
            "aging_dispos": aging_dispos,
            "stopclock": stopclock,
            "aging_close": aging_close,
            "aging_berjalan": aging_berjalan,
            "durasi_hari": durasi_hari,
            "aging_category": aging_category,
            "aging_label": aging_label,
            "update_pekerjaan": update_pekerjaan,
            "bulan": bulan,
            "tahun": tahun,
            "bulan_finish": bulan_finish,
            "tahun_finish": tahun_finish
        })

    print(f"[SUCCESS] Berhasil mengekstrak {len(records)} baris Master RAW data.")
    return records

# ==============================================================================
# 2. EXTRACTION: GANGGUAN SHEET
# ==============================================================================
def extract_gangguan(wb):
    print("[INFO] Mengekstrak sheet: GANGGUAN...")
    if "GANGGUAN" not in wb.sheetnames:
        print("[WARN] Sheet GANGGUAN tidak ditemukan!")
        return []

    ws = wb["GANGGUAN"]
    rows = list(ws.iter_rows(values_only=True))
    data_rows = rows[1:]

    gangguan_records = []
    for r_idx, row in enumerate(data_rows, start=2):
        if not any(row):
            continue

        no_val = row[0] if len(row) > 0 else None
        pa_tiket = clean_text(row[3] if len(row) > 3 else "")
        terminating = clean_text(row[4] if len(row) > 4 else "")
        jenis_gangguan = clean_text(row[8] if len(row) > 8 else "")

        if no_val is None and not pa_tiket and not terminating and not jenis_gangguan:
            continue

        g_id = len(gangguan_records) + 1
        tgl_dispos = parse_date(row[1] if len(row) > 1 else None)
        tgl_selesai = parse_date(row[2] if len(row) > 2 else None)
        pic_disposisi = clean_text(row[5] if len(row) > 5 else "")
        pic_perijinan = clean_text(row[6] if len(row) > 6 else "")
        pic_admin = clean_text(row[7] if len(row) > 7 else "")

        raw_biaya_awal = row[9] if len(row) > 9 else None
        biaya_awal, _ = parse_currency(raw_biaya_awal)

        raw_biaya_gangguan = row[10] if len(row) > 10 else None
        biaya_gangguan, _ = parse_currency(raw_biaya_gangguan)

        update_gangguan = clean_text(row[11] if len(row) > 11 else "")

        # Status tiket: Jika ada tanggal selesai atau update menyebutkan selesai -> Selesai, selain itu -> Aktif (Open)
        is_selesai = 1 if (tgl_selesai is not None or "selesai" in update_gangguan.lower()) else 0
        status_tiket = "Selesai" if is_selesai == 1 else "Aktif (Open)"

        # Hitung SLA durasi hari penanganan
        durasi_hari = 0
        if tgl_dispos and tgl_selesai:
            try:
                d1 = datetime.datetime.strptime(tgl_dispos, "%Y-%m-%d")
                d2 = datetime.datetime.strptime(tgl_selesai, "%Y-%m-%d")
                durasi_hari = max(0, (d2 - d1).days)
            except Exception:
                pass

        # Resolusi koordinat untuk peta
        lat, lng, coord_type = resolve_coordinates("", terminating, g_id + 1000)

        gangguan_records.append({
            "id": g_id,
            "no_tiket": pa_tiket or f"TIKET-{g_id:04d}",
            "tgl_dispos": tgl_dispos,
            "tgl_selesai": tgl_selesai,
            "terminating": terminating,
            "latitude": lat,
            "longitude": lng,
            "coord_type": coord_type,
            "pic_disposisi": pic_disposisi,
            "pic_perijinan": pic_perijinan,
            "pic_admin": pic_admin,
            "jenis_gangguan": jenis_gangguan,
            "biaya_awal": biaya_awal,
            "biaya_gangguan": biaya_gangguan,
            "status": status_tiket,
            "is_selesai": is_selesai,
            "durasi_hari": durasi_hari,
            "update_gangguan": update_gangguan
        })

    print(f"[SUCCESS] Berhasil mengekstrak {len(gangguan_records)} tiket gangguan.")
    return gangguan_records

# ==============================================================================
# 3. EXTRACTION: REKAP EFISIENSI BIAYA PERIJINAN
# ==============================================================================
def extract_rekap_efisiensi(wb):
    print("[INFO] Mengekstrak sheet: Rekap Efesiensi Biaya Perijinan...")
    if "Rekap Efesiensi Biaya Perijinan" not in wb.sheetnames:
        return []

    ws = wb["Rekap Efesiensi Biaya Perijinan"]
    rekap_list = []
    
    # Cari baris yang memuat Tahun 2024, 2025, 2026
    for row in ws.iter_rows(values_only=True):
        if not any(row):
            continue
        first_val = str(row[0]).strip() if row[0] is not None else ""
        if first_val in ("2024", "2025", "2026"):
            tahun = int(first_val)
            total_disposisi = parse_int_safe(row[1], 0)
            ada_biaya = parse_int_safe(row[2], 0)
            nilai_awal = parse_currency(row[3])[0]
            nilai_akhir = parse_currency(row[4])[0]
            tidak_ada_biaya = parse_int_safe(row[5], 0)
            efisiensi_rp = max(0, nilai_awal - nilai_akhir) if nilai_awal > 0 else 0
            efisiensi_pct = round((efisiensi_rp / nilai_awal * 100), 2) if nilai_awal > 0 else 0.0

            rekap_list.append({
                "tahun": tahun,
                "total_disposisi": total_disposisi,
                "ada_biaya": ada_biaya,
                "tidak_ada_biaya": tidak_ada_biaya,
                "nilai_awal": nilai_awal,
                "nilai_akhir": nilai_akhir,
                "efisiensi_rupiah": efisiensi_rp,
                "efisiensi_persen": efisiensi_pct
            })

    print(f"[SUCCESS] Berhasil mengekstrak {len(rekap_list)} baris Rekap Efisiensi Biaya.")
    return rekap_list

# ==============================================================================
# 4. WORKLOAD & KPI SUMMARY AGGREGATOR
# ==============================================================================
def generate_master_summary(records, gangguan_records, rekap_efisiensi):
    total_records = len(records)
    total_biaya_awal = sum(r["biaya_permintaan_awal"] for r in records)
    total_biaya_final = sum(r["biaya_final"] for r in records)
    total_efisiensi = sum(r["efisiensi_rupiah"] for r in records)
    persen_efisiensi_total = round((total_efisiensi / total_biaya_awal * 100), 2) if total_biaya_awal > 0 else 0
    
    # Status breakdown
    status_counts = {"Finish": 0, "Ongoing": 0, "Hold": 0, "Cancel": 0}
    for r in records:
        st = r["progress"]
        status_counts[st] = status_counts.get(st, 0) + 1

    status_pct = {
        k: round((v / total_records * 100), 1) if total_records > 0 else 0 
        for k, v in status_counts.items()
    }

    # Zero cost vs Paid
    zero_cost_count = sum(1 for r in records if r["is_berbiaya"] == 0)
    paid_count = total_records - zero_cost_count
    zero_cost_pct = round((zero_cost_count / total_records * 100), 1) if total_records > 0 else 0

    # Summary Gangguan
    total_gangguan = len(gangguan_records)
    gangguan_aktif = sum(1 for g in gangguan_records if g["is_selesai"] == 0)
    gangguan_selesai = total_gangguan - gangguan_aktif
    gangguan_resolution_rate = round((gangguan_selesai / total_gangguan * 100), 1) if total_gangguan > 0 else 0
    total_biaya_gangguan = sum(g["biaya_gangguan"] for g in gangguan_records)

    # Workload per PIC
    pic_matrix = {}
    for r in records:
        pic = r["pic_perijinan"] or "Unassigned"
        if pic not in pic_matrix:
            pic_matrix[pic] = {
                "pic": pic,
                "total_penugasan": 0,
                "finish": 0,
                "ongoing": 0,
                "cancel": 0,
                "hold": 0,
                "biaya_awal": 0,
                "biaya_final": 0,
                "efisiensi_rupiah": 0,
                "sla_days": [],
                "gangguan_count": 0
            }
        pic_matrix[pic]["total_penugasan"] += 1
        st = r["progress"].lower()
        if st in pic_matrix[pic]:
            pic_matrix[pic][st] += 1
        pic_matrix[pic]["biaya_awal"] += r["biaya_permintaan_awal"]
        pic_matrix[pic]["biaya_final"] += r["biaya_final"]
        pic_matrix[pic]["efisiensi_rupiah"] += r["efisiensi_rupiah"]
        if r["progress"] == "Finish" and r["durasi_hari"] is not None and r["durasi_hari"] >= 0:
            pic_matrix[pic]["sla_days"].append(r["durasi_hari"])

    # Tambahkan hitungan gangguan per PIC
    for g in gangguan_records:
        pic_g = g["pic_perijinan"]
        if not pic_g: continue
        # Cocokkan nama PIC
        matched = False
        for pic in pic_matrix:
            if pic_g.lower() in pic.lower() or pic.lower() in pic_g.lower():
                pic_matrix[pic]["gangguan_count"] += 1
                matched = True
                break

    # Finalisasi metrik per PIC
    workload_list = []
    for pic, d in pic_matrix.items():
        finish_rate = round((d["finish"] / d["total_penugasan"] * 100), 1) if d["total_penugasan"] > 0 else 0
        efisiensi_pct = round((d["efisiensi_rupiah"] / d["biaya_awal"] * 100), 1) if d["biaya_awal"] > 0 else 0
        avg_sla = round(sum(d["sla_days"]) / len(d["sla_days"]), 1) if d["sla_days"] else 0
        workload_list.append({
            "pic": pic,
            "total_penugasan": d["total_penugasan"],
            "finish": d["finish"],
            "ongoing": d["ongoing"],
            "cancel": d["cancel"],
            "hold": d["hold"],
            "finish_rate": finish_rate,
            "biaya_awal": d["biaya_awal"],
            "biaya_final": d["biaya_final"],
            "efisiensi_rupiah": d["efisiensi_rupiah"],
            "efisiensi_persen": efisiensi_pct,
            "avg_sla_hari": avg_sla,
            "gangguan_count": d["gangguan_count"]
        })

    workload_list.sort(key=lambda x: x["total_penugasan"], reverse=True)

    # Rata-rata SLA keseluruhan
    selesai_records = [r["durasi_hari"] for r in records if r["progress"] == "Finish" and r["durasi_hari"] is not None and r["durasi_hari"] >= 0]
    avg_sla = round(sum(selesai_records) / len(selesai_records), 1) if selesai_records else 0

    return {
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_proyek": total_records,
        "total_biaya_awal": total_biaya_awal,
        "total_biaya_final": total_biaya_final,
        "total_efisiensi_rupiah": total_efisiensi,
        "total_efisiensi_persen": persen_efisiensi_total,
        "status_distribution": status_counts,
        "status_percentage": status_pct,
        "zero_cost_count": zero_cost_count,
        "zero_cost_percent": zero_cost_pct,
        "paid_count": paid_count,
        "avg_sla_hari": avg_sla,
        "gangguan_summary": {
            "total_tiket": total_gangguan,
            "aktif": gangguan_aktif,
            "selesai": gangguan_selesai,
            "resolution_rate": gangguan_resolution_rate,
            "total_biaya": total_biaya_gangguan
        },
        "workload_performa": workload_list,
        "rekap_efisiensi": rekap_efisiensi
    }

# ==============================================================================
# 5. SAVING OUTPUTS (JSON, JS, SQLITE)
# ==============================================================================
def save_outputs(records, gangguan_records, rekap_efisiensi, summary):
    payload = {
        "metadata": {
            "source_file": EXCEL_FILE,
            "version": "2.0",
            "last_updated": summary["generated_at"],
            "total_master_records": len(records),
            "total_gangguan_records": len(gangguan_records),
            "total_rekap_records": len(rekap_efisiensi)
        },
        "summary": summary,
        "data": records,
        "gangguan": gangguan_records,
        "rekap_efisiensi": rekap_efisiensi
    }

    # 1. JSON
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"[SUCCESS] File JSON berhasil disimpan: {OUTPUT_JSON} ({os.path.getsize(OUTPUT_JSON):,} bytes)")

    # 2. Standalone JS
    with open(OUTPUT_JS, "w", encoding="utf-8") as f:
        f.write("window.SITAC_RAW_DATA = ")
        json.dump(payload, f, ensure_ascii=False)
        f.write(";\n")
    print(f"[SUCCESS] File JS berhasil disimpan: {OUTPUT_JS} ({os.path.getsize(OUTPUT_JS):,} bytes)")

    # 3. SQLite Database
    if os.path.exists(OUTPUT_DB):
        try:
            os.remove(OUTPUT_DB)
        except Exception:
            pass

    conn = sqlite3.connect(OUTPUT_DB)
    cur = conn.cursor()

    # Tabel Master Records
    cur.execute("""
    CREATE TABLE sitac_records (
        id INTEGER PRIMARY KEY,
        excel_no INTEGER,
        no_pa TEXT,
        pelanggan TEXT,
        ptl TEXT,
        pic_perijinan TEXT,
        tiering TEXT,
        poin INTEGER,
        terminating TEXT,
        koordinat_raw TEXT,
        latitude REAL,
        longitude REAL,
        coord_type TEXT,
        biaya_permintaan_awal INTEGER,
        catatan_biaya_awal TEXT,
        sewa_otc TEXT,
        biaya_final INTEGER,
        biaya_sewa_bulan INTEGER,
        efisiensi_rupiah INTEGER,
        efisiensi_persen REAL,
        is_berbiaya INTEGER,
        progress TEXT,
        date_pa TEXT,
        date_dispos TEXT,
        date_close TEXT,
        aging_dispos INTEGER,
        stopclock INTEGER,
        aging_close INTEGER,
        aging_berjalan INTEGER,
        durasi_hari INTEGER,
        aging_category TEXT,
        aging_label TEXT,
        update_pekerjaan TEXT,
        bulan INTEGER,
        tahun INTEGER,
        bulan_finish INTEGER,
        tahun_finish INTEGER
    );
    """)

    cur.execute("CREATE INDEX idx_status ON sitac_records(progress);")
    cur.execute("CREATE INDEX idx_pic ON sitac_records(pic_perijinan);")
    cur.execute("CREATE INDEX idx_tahun ON sitac_records(tahun);")
    cur.execute("CREATE INDEX idx_biaya_final ON sitac_records(biaya_final);")
    cur.execute("CREATE INDEX idx_no_pa ON sitac_records(no_pa);")

    insert_master_sql = """
    INSERT INTO sitac_records VALUES (
        :id, :excel_no, :no_pa, :pelanggan, :ptl, :pic_perijinan, :tiering, :poin, :terminating,
        :koordinat_raw, :latitude, :longitude, :coord_type, :biaya_permintaan_awal,
        :catatan_biaya_awal, :sewa_otc, :biaya_final, :biaya_sewa_bulan, :efisiensi_rupiah,
        :efisiensi_persen, :is_berbiaya, :progress, :date_pa, :date_dispos, :date_close,
        :aging_dispos, :stopclock, :aging_close, :aging_berjalan, :durasi_hari,
        :aging_category, :aging_label, :update_pekerjaan, :bulan, :tahun, :bulan_finish,
        :tahun_finish
    );
    """
    cur.executemany(insert_master_sql, records)

    # Tabel Gangguan
    cur.execute("""
    CREATE TABLE gangguan_records (
        id INTEGER PRIMARY KEY,
        no_tiket TEXT,
        tgl_dispos TEXT,
        tgl_selesai TEXT,
        terminating TEXT,
        latitude REAL,
        longitude REAL,
        coord_type TEXT,
        pic_disposisi TEXT,
        pic_perijinan TEXT,
        pic_admin TEXT,
        jenis_gangguan TEXT,
        biaya_awal INTEGER,
        biaya_gangguan INTEGER,
        status TEXT,
        is_selesai INTEGER,
        durasi_hari INTEGER,
        update_gangguan TEXT
    );
    """)
    cur.execute("CREATE INDEX idx_gangguan_status ON gangguan_records(status);")
    cur.execute("CREATE INDEX idx_gangguan_pic ON gangguan_records(pic_perijinan);")

    insert_gangguan_sql = """
    INSERT INTO gangguan_records VALUES (
        :id, :no_tiket, :tgl_dispos, :tgl_selesai, :terminating, :latitude, :longitude,
        :coord_type, :pic_disposisi, :pic_perijinan, :pic_admin, :jenis_gangguan,
        :biaya_awal, :biaya_gangguan, :status, :is_selesai, :durasi_hari, :update_gangguan
    );
    """
    cur.executemany(insert_gangguan_sql, gangguan_records)

    # Tabel Rekap Efisiensi
    cur.execute("""
    CREATE TABLE rekap_efisiensi (
        tahun INTEGER PRIMARY KEY,
        total_disposisi INTEGER,
        ada_biaya INTEGER,
        tidak_ada_biaya INTEGER,
        nilai_awal INTEGER,
        nilai_akhir INTEGER,
        efisiensi_rupiah INTEGER,
        efisiensi_persen REAL
    );
    """)
    insert_rekap_sql = """
    INSERT INTO rekap_efisiensi VALUES (
        :tahun, :total_disposisi, :ada_biaya, :tidak_ada_biaya, :nilai_awal, :nilai_akhir,
        :efisiensi_rupiah, :efisiensi_persen
    );
    """
    cur.executemany(insert_rekap_sql, rekap_efisiensi)

    # Views Analitis
    cur.execute("""
    CREATE VIEW v_kpi_per_pic AS
    SELECT 
        pic_perijinan,
        COUNT(*) as total_proyek,
        SUM(CASE WHEN progress = 'Finish' THEN 1 ELSE 0 END) as finish_count,
        ROUND(CAST(SUM(CASE WHEN progress = 'Finish' THEN 1 ELSE 0 END) AS REAL) / COUNT(*) * 100, 1) as finish_rate,
        SUM(CASE WHEN progress = 'Ongoing' THEN 1 ELSE 0 END) as ongoing_count,
        SUM(CASE WHEN progress = 'Cancel' THEN 1 ELSE 0 END) as cancel_count,
        SUM(CASE WHEN progress = 'Hold' THEN 1 ELSE 0 END) as hold_count,
        SUM(biaya_permintaan_awal) as total_biaya_awal,
        SUM(biaya_final) as total_biaya_final,
        SUM(efisiensi_rupiah) as total_efisiensi_rupiah,
        ROUND(AVG(CASE WHEN progress = 'Finish' AND durasi_hari >= 0 THEN durasi_hari END), 1) as avg_sla_hari
    FROM sitac_records
    GROUP BY pic_perijinan
    ORDER BY total_proyek DESC;
    """)

    conn.commit()
    conn.close()
    print(f"[SUCCESS] Database SQLite berhasil disimpan: {OUTPUT_DB} ({os.path.getsize(OUTPUT_DB):,} bytes)")

def main():
    print("=" * 65)
    print(" MULTI-SHEET DATA PIPELINE SITAC & GANGGUAN 2026")
    print("=" * 65)
    if not os.path.exists(EXCEL_FILE):
        raise FileNotFoundError(f"File {EXCEL_FILE} tidak ditemukan!")

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    master_records = extract_master_raw_data(wb)
    gangguan_records = extract_gangguan(wb)
    rekap_efisiensi = extract_rekap_efisiensi(wb)
    
    summary = generate_master_summary(master_records, gangguan_records, rekap_efisiensi)
    save_outputs(master_records, gangguan_records, rekap_efisiensi, summary)

    print("\n--- RINGKASAN METRIK MASTER ---")
    print(f"Total PA           : {summary['total_proyek']}")
    print(f"Status Breakdown   : {summary['status_distribution']} ({summary['status_percentage']})")
    print(f"Total Biaya Awal   : Rp {summary['total_biaya_awal']:,}")
    print(f"Total Biaya Final  : Rp {summary['total_biaya_final']:,}")
    print(f"Total Penghematan  : Rp {summary['total_efisiensi_rupiah']:,} ({summary['total_efisiensi_persen']}%)")
    print(f"Zero Cost Projects : {summary['zero_cost_count']} PA ({summary['zero_cost_percent']}%)")
    print(f"Tiket Gangguan     : {summary['gangguan_summary']['total_tiket']} Total | {summary['gangguan_summary']['aktif']} Aktif | {summary['gangguan_summary']['selesai']} Selesai ({summary['gangguan_summary']['resolution_rate']}%)")
    print("=" * 65)

if __name__ == "__main__":
    main()
